from pathlib import Path

from openpyxl import load_workbook

from app.domain.models import (
    Choice,
    Constraint,
    Form,
    Question,
    QuestionType,
    RelevanceCondition,
)


class XLSFormParser:
    REQUIRED_SHEETS = {"survey", "choices"}

    QUESTION_TYPES = {
        "text": QuestionType.TEXT,
        "integer": QuestionType.INTEGER,
        "decimal": QuestionType.DECIMAL,
        "boolean": QuestionType.BOOLEAN,
        "date": QuestionType.DATE,
        "time": QuestionType.TIME,
    }

    SKIPPED_TYPES = {
        "begin_group",
        "end_group",
        "begin_repeat",
        "end_repeat",
        "start",
        "end",
        "today",
        "deviceid",
    }

    def parse(self, file_path: str | Path) -> Form:
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(
                f"XLSForm file not found: {file_path}"
            )

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        try:
            missing_sheets = (
                self.REQUIRED_SHEETS
                - set(workbook.sheetnames)
            )

            if missing_sheets:
                raise ValueError(
                    f"XLSForm is missing required sheets: "
                    f"{sorted(missing_sheets)}"
                )

            settings = self._parse_settings(workbook)
            choices = self._parse_choices(workbook)
            questions = self._parse_questions(
                workbook,
                choices,
            )

            return Form(
                id=settings.get(
                    "form_id",
                    file_path.stem,
                ),
                title=settings.get(
                    "form_title",
                    file_path.stem,
                ),
                questions=questions,
                version=settings.get("version"),
                metadata=settings,
            )
        finally:
            workbook.close()

    def _parse_settings(self, workbook) -> dict[str, str]:
        if "settings" not in workbook.sheetnames:
            return {}

        worksheet = workbook["settings"]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return {}

        headers = [
            str(header).strip() if header is not None else ""
            for header in headers
        ]

        settings = {}

        for row in rows:
            for header, value in zip(headers, row):
                if header and value is not None:
                    settings[header] = str(value).strip()

        return settings

    def _parse_choices(
        self,
        workbook,
    ) -> dict[str, list[Choice]]:
        worksheet = workbook["choices"]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return {}

        headers = [
            str(header).strip() if header is not None else ""
            for header in headers
        ]

        choices = {}

        for row in rows:
            data = dict(zip(headers, row))

            list_name = str(
                data.get("list_name") or ""
            ).strip()

            value = str(
                data.get("name") or ""
            ).strip()

            label = str(
                data.get("label") or ""
            ).strip()

            if not list_name or not value:
                continue

            choices.setdefault(list_name, []).append(
                Choice(
                    value=value,
                    label=label,
                )
            )

        return choices

    def _parse_questions(
        self,
        workbook,
        choices: dict[str, list[Choice]],
    ) -> list[Question]:
        worksheet = workbook["survey"]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [
            str(header).strip() if header is not None else ""
            for header in headers
        ]

        questions = []

        for row in rows:
            data = dict(zip(headers, row))

            question_type = str(
                data.get("type") or ""
            ).strip()

            name = str(
                data.get("name") or ""
            ).strip()

            label = str(
                data.get("label") or ""
            ).strip()

            if (
                not question_type
                or question_type in self.SKIPPED_TYPES
            ):
                continue

            if not name:
                continue

            mapped_type = self._map_question_type(
                question_type
            )

            if mapped_type is None:
                continue

            question_choices = []

            if mapped_type in {
                QuestionType.SELECT_ONE,
                QuestionType.SELECT_MULTIPLE,
            }:
                list_name = self._get_choice_list_name(
                    question_type
                )

                question_choices = choices.get(
                    list_name,
                    [],
                )

            question = Question(
                name=name,
                type=mapped_type,
                label=label,
                required=self._parse_required(
                    data.get("required")
                ),
                choices=question_choices,
                relevance=self._parse_relevance(
                    data.get("relevant")
                ),
                constraint=self._parse_constraint(
                    data.get("constraint"),
                    data.get("constraint_message"),
                ),
            )

            questions.append(question)

        return questions

    def _map_question_type(
        self,
        question_type: str,
    ) -> QuestionType | None:
        if question_type in self.QUESTION_TYPES:
            return self.QUESTION_TYPES[question_type]

        if question_type == "yes_no":
            return QuestionType.BOOLEAN

        if question_type.startswith("select_one"):
            return QuestionType.SELECT_ONE

        if question_type.startswith("select_multiple"):
            return QuestionType.SELECT_MULTIPLE

        return None

    def _get_choice_list_name(
        self,
        question_type: str,
    ) -> str:
        return question_type.split()[1]

    def _parse_required(
        self,
        value,
    ) -> bool:
        if value is None:
            return False

        return str(value).strip().lower() in {
            "yes",
            "true",
            "1",
        }

    def _parse_relevance(
        self,
        value,
    ) -> RelevanceCondition | None:
        if value is None or not str(value).strip():
            return None

        return RelevanceCondition(
            expression=str(value).strip()
        )

    def _parse_constraint(
        self,
        expression,
        message,
    ) -> Constraint | None:
        if (
            expression is None
            or not str(expression).strip()
        ):
            return None

        constraint_message = (
            str(message).strip()
            if message is not None
            else None
        )

        return Constraint(
            expression=str(expression).strip(),
            message=constraint_message,
        )